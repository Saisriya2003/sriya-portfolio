"""Generate an ATS-safe single-column resume (DOCX)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


OUTPUTS = [
    Path(r"C:\Users\saisr\OneDrive\Desktop\personal_files\Pettem_Sai_Sriya_Resume.docx"),
    Path(r"C:\Users\saisr\sriya-portfolio\Pettem_Sai_Sriya_Resume.docx"),
]


def set_run_font(run, name="Calibri", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=0, line=240):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Twips(line)


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=7, after=3, line=230)
    add_text(p, text.upper(), size=11.5, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    add_bottom_border(p)
    return p


def body(doc, text, after=3):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=after, line=216)
    add_text(p, text, size=10.5)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=1, line=226)
    p.clear()
    add_text(p, text, size=10.5)
    return p


def role_line(doc, left, right):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=0, line=216)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.5), WD_TAB_ALIGNMENT.RIGHT)
    add_text(p, left, size=10.5, bold=True)
    add_text(p, "\t" + right, size=10.5)
    return p


def sub_line(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=220)
    add_text(p, text, size=10)
    return p


def build():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    bullet_style = styles["List Bullet"]
    bullet_style.font.name = "Calibri"
    bullet_style.font.size = Pt(10.5)
    pf = bullet_style.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(-0.15)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, before=0, after=0, line=260)
    add_text(name, "PETTEM SAI SRIYA", size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, before=0, after=1, line=210)
    add_text(title, "Full Stack Developer  |  Python  ·  React  ·  Node.js  ·  GenAI", size=11, bold=True)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact, before=0, after=2, line=210)
    add_text(
        contact,
        "Hyderabad, Telangana, India  |  9921890019  |  saisriyavarma@gmail.com  |  linkedin.com/in/saisriyavarma  |  github.com/Saisriya2003",
        size=10,
    )

    heading(doc, "Professional Summary")
    body(
        doc,
        "Full Stack Developer (immediately available, Hyderabad) with 1+ year building React, Node.js, "
        "and Python products. Shipped REST APIs, dashboards, and Generative AI features including RAG, "
        "embeddings, semantic search, and LLM document Q&A. B.Tech ECE, CGPA 7.6.",
        after=1,
    )

    heading(doc, "Technical Skills")
    skills = [
        ("Languages: ", "Python, JavaScript, TypeScript, SQL, HTML, CSS"),
        ("Frameworks: ", "React, FastAPI, Node.js, Express.js, Socket.IO, Tailwind CSS"),
        ("Databases: ", "PostgreSQL, MongoDB, MySQL, Redis, SQLAlchemy"),
        ("Developer Tools: ", "Git, GitHub Actions, Docker, Pytest, Vite, REST APIs, JWT"),
        (
            "AI / ML: ",
            "Generative AI, RAG, Large Language Models (LLM), Embeddings, Semantic Search, Neural Networks, Voice AI",
        ),
    ]
    for label, value in skills:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=1, line=226)
        add_text(p, label, size=10.5, bold=True)
        add_text(p, value, size=10.5)

    heading(doc, "Professional Experience")

    role_line(
        doc,
        "AI Evaluation Specialist  |  Handshake AI",
        "Jul 2026 – Present",
    )
    for item in [
        "Design terminal-based benchmark tasks that evaluate coding agents on multi-step Linux, Git, Bash, and Python workflows.",
        "Build reproducible environments with automated validation so tasks score consistently across model runs and reviewers.",
        "Revise task specifications and graders after human review when an agent passes a check but misses the intended engineering behavior.",
    ]:
        bullet(doc, item)

    role_line(
        doc,
        "Full Stack Developer  |  GenAI Lakes  |  Hyderabad, India",
        "Jul 2025 – Sep 2026",
    )
    for item in [
        "Designed and developed scalable full-stack web applications using modern frontend and backend technologies.",
        "Built responsive, user-centric interfaces while integrating secure RESTful APIs and backend services.",
        "Optimized application performance, database operations, and system reliability through debugging and testing.",
        "Collaborated with cross-functional teams to deliver AI-driven and web-based solutions aligned with business requirements.",
    ]:
        bullet(doc, item)

    role_line(
        doc,
        "LLM Prompt Engineer  |  Turing",
        "Aug 2025 – Jan 2026",
    )
    for item in [
        "Delivered high-quality multi-turn AI conversations focused on realistic user–assistant interactions and contextual accuracy.",
        "Designed and optimized prompt workflows to improve reasoning, instruction-following, and response consistency.",
        "Evaluated AI-generated outputs using quality guidelines for relevance, clarity, and conversational flow.",
    ]:
        bullet(doc, item)

    role_line(
        doc,
        "Aether  |  Outlier.ai",
        "May 2026 – Jun 2026",
    )
    for item in [
        "Recorded multilingual voice samples using predefined prompts to support AI model training.",
        "Followed strict quality and language guidelines to ensure accurate and consistent speech data.",
    ]:
        bullet(doc, item)

    heading(doc, "Projects")

    role_line(
        doc,
        "PulseBoard — Real-Time Team Collaboration Platform  |  github.com/Saisriya2003/pulseboard",
        "September 2026",
    )
    sub_line(doc, "Tech: Python, FastAPI, Socket.IO, React, TypeScript, PostgreSQL, Redis, Docker, Pytest, GitHub Actions")
    for item in [
        "Built a real-time kanban platform: workspaces, drag-and-drop boards, labels, checklists, @mentions, reactions, and live presence over Socket.IO; 30 Pytest integration tests run in GitHub Actions CI.",
        "Designed a four-role permission model (owner/admin/member/viewer) enforced server-side with JWT auth, invitations, audit logging, and Redis caching; containerized with Docker Compose.",
    ]:
        bullet(doc, item)

    role_line(
        doc,
        "PostPilot — AI Social Media Management Platform",
        "July 2026 – August 2026",
    )
    sub_line(doc, "Tech: TypeScript, React, Python, FastAPI, PostgreSQL, MongoDB, REST APIs, Tailwind CSS")
    bullet(
        doc,
        "Built a full-stack platform to create, schedule, and publish content to YouTube, Instagram, Facebook, Twitter, LinkedIn, and TikTok: TypeScript React frontend, FastAPI backend with JWT, PostgreSQL, MongoDB, OAuth, Gemini captions, and analytics dashboards.",
    )

    role_line(
        doc,
        "Manager Monitoring Assistant — AI Status Call System",
        "2025",
    )
    bullet(
        doc,
        "Shipped an AI status-call system for employee and field-agent tracking with automated outbound calls, transcription, live dashboards, alerts, and reporting. Tech: React, Node.js, Express.js, MongoDB, MySQL, Voice AI.",
    )

    role_line(
        doc,
        "AI-Based RAG Application — Document Question Answering",
        "2025",
    )
    bullet(
        doc,
        "Built a Retrieval-Augmented Generation pipeline with document ingestion, embeddings, semantic search, and context-aware LLM answers. Tech: Python, FastAPI, React, PostgreSQL.",
    )

    role_line(
        doc,
        "Identification of Fake Profiles Across Online Social Networks  (Team)",
        "April 2025 – May 2025",
    )
    bullet(
        doc,
        "Built a Python deep learning model using artificial neural networks and backpropagation to score friend-request authenticity on social-network data.",
    )

    heading(doc, "Education")
    body(
        doc,
        "B.Tech, Electronics and Communication Engineering  |  Malla Reddy Institute of Technology and Science, Hyderabad  |  Nov 2021 – Aug 2025  |  CGPA 7.6",
        after=1,
    )
    body(
        doc,
        "Intermediate  |  Telangana State Residential Junior College (Girls), Huzurabad  |  Jun 2019 – Jul 2021  |  88%",
        after=1,
    )
    body(
        doc,
        "Secondary School  |  Trinity High School, Mancherial  |  May 2018 – Jun 2019  |  GPA 9.2",
        after=1,
    )

    heading(doc, "Certifications")
    body(
        doc,
        "AWS and Google Cloud foundational badges. Certifications in Generative AI, Prompt Engineering, Machine Learning, React.js, JavaScript, and Python. AICTE Frontend Development Internship; API Integration, Database Management, and Full-Stack Development.",
        after=0,
    )

    saved = []
    for path in OUTPUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        saved.append(str(path))
    return saved


if __name__ == "__main__":
    for p in build():
        print(p)
